from __future__ import annotations

import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from .models import FieldDefinition, PipelineStep, SourceDefinition


SOURCE_ID_RE = re.compile(r"\bS\d{2}\b", re.IGNORECASE)


def clean_cell(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def norm(value: Any) -> str:
    text = clean_cell(value) or ""
    text = unicodedata.normalize("NFKC", text).casefold().replace("ё", "е")
    return " ".join(text.split())


def _table_rows(sheet: Any) -> list[dict[str, str | None]]:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    header_index = next(
        (index for index, row in enumerate(rows) if any(clean_cell(value) for value in row)),
        None,
    )
    if header_index is None:
        return []
    raw_headers = rows[header_index]
    headers: list[str] = []
    for index, value in enumerate(raw_headers, start=1):
        headers.append(clean_cell(value) or f"column_{index}")

    result: list[dict[str, str | None]] = []
    for row in rows[header_index + 1 :]:
        values = list(row) + [None] * max(0, len(headers) - len(row))
        item = {headers[index]: clean_cell(values[index]) for index in range(len(headers))}
        if any(value is not None for value in item.values()):
            result.append(item)
    return result


class SourceMap:
    """Загружает все листы схемы Andromeda из Excel без захардкоживания строк."""

    def __init__(
        self,
        sources: list[SourceDefinition],
        fields: list[FieldDefinition],
        pipeline: list[PipelineStep],
        risks: list[dict[str, str | None]],
        workbook_path: Path,
    ) -> None:
        self.sources = sources
        self.fields = fields
        self.pipeline = pipeline
        self.risks = risks
        self.workbook_path = workbook_path
        self._by_id = {source.id: source for source in sources}
        self._aliases: dict[str, str] = {}
        for source in sources:
            for alias in (source.id, source.name, source.url):
                self._aliases[norm(alias)] = source.id

    @classmethod
    def from_xlsx(cls, path: str | Path) -> "SourceMap":
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - dependency is declared in pyproject
            raise RuntimeError("Для чтения source map нужен пакет openpyxl") from exc

        workbook_path = Path(path).expanduser().resolve()
        if not workbook_path.exists():
            raise FileNotFoundError(f"Source map не найден: {workbook_path}")

        workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
        sheets = {_sheet_key(name): workbook[name] for name in workbook.sheetnames}
        source_rows = _table_rows(_required_sheet(sheets, "источники"))
        field_rows = _table_rows(_required_sheet(sheets, "карта данных"))
        pipeline_rows = _table_rows(_required_sheet(sheets, "пайплайн"))
        risk_sheet = sheets.get(_sheet_key("пробелы и риски"))
        risk_rows = _table_rows(risk_sheet) if risk_sheet is not None else []

        sources = [
            SourceDefinition(
                id=clean_cell(row.get("ID")) or "",
                name=clean_cell(row.get("Источник")) or "",
                url=clean_cell(row.get("URL")) or "",
                officiality=clean_cell(row.get("Официальность")),
                scope=clean_cell(row.get("Охват")),
                format=clean_cell(row.get("Формат")),
                availability=clean_cell(row.get("Доступность")),
                method=clean_cell(row.get("Метод")),
                data_description=clean_cell(row.get("Какие данные")),
                history=clean_cell(row.get("История")),
                refresh=clean_cell(row.get("Обновление")),
                note=clean_cell(row.get("Примечание")),
            )
            for row in source_rows
            if clean_cell(row.get("ID"))
        ]

        source_aliases: dict[str, str] = {}
        for source in sources:
            for alias in (source.id, source.name, source.url):
                source_aliases[norm(alias)] = source.id

        fields: list[FieldDefinition] = []
        for row in field_rows:
            raw_id = clean_cell(row.get("ID"))
            if not raw_id or not raw_id.isdigit():
                continue
            primary_text = clean_cell(row.get("Основной источник"))
            reserve_text = clean_cell(row.get("Резерв / сверка"))
            fields.append(
                FieldDefinition(
                    id=int(raw_id),
                    block=clean_cell(row.get("Блок")) or "",
                    name=clean_cell(row.get("Поле Andromeda")) or "",
                    entity=clean_cell(row.get("Сущность")) or "",
                    acquisition_status=clean_cell(row.get("Статус получения")) or "",
                    primary_source=primary_text,
                    url=clean_cell(row.get("URL")),
                    extraction_target=clean_cell(row.get("Что именно получать")),
                    format=clean_cell(row.get("Формат")),
                    recommended_method=clean_cell(row.get("Рекомендуемый метод")),
                    coverage=clean_cell(row.get("Покрытие")),
                    reserve=reserve_text,
                    reserve_url=clean_cell(row.get("URL резерва")),
                    history=clean_cell(row.get("История")),
                    refresh=clean_cell(row.get("Частота обновления")),
                    reliability=clean_cell(row.get("Надёжность")),
                    priority=clean_cell(row.get("Приоритет")),
                    limitation=clean_cell(row.get("Честное ограничение / комментарий")),
                    source_ids=resolve_source_ids(primary_text, source_aliases),
                    reserve_source_ids=resolve_source_ids(reserve_text, source_aliases),
                )
            )

        pipeline: list[PipelineStep] = []
        for row in pipeline_rows:
            raw_order = clean_cell(row.get("Порядок"))
            if not raw_order or not raw_order.isdigit():
                continue
            pipeline.append(
                PipelineStep(
                    order=int(raw_order),
                    module=clean_cell(row.get("Модуль")) or "",
                    sources=clean_cell(row.get("Источники")),
                    output=clean_cell(row.get("Что получает")),
                    technology=clean_cell(row.get("Технология")),
                    depends_on=clean_cell(row.get("Зависит от")),
                    frequency=clean_cell(row.get("Частота")),
                    fail_safe=clean_cell(row.get("Fail-safe / правило приоритета")),
                )
            )

        result = cls(sources, fields, pipeline, risk_rows, workbook_path)
        result.validate()
        return result

    def source(self, source_id: str) -> SourceDefinition:
        return self._by_id[source_id]

    def resolve_source_ids(self, value: str | None) -> tuple[str, ...]:
        return resolve_source_ids(value, self._aliases)

    def selected_sources(self, requested: Iterable[str] | None = None) -> list[SourceDefinition]:
        if not requested:
            return list(self.sources)
        wanted: set[str] = set()
        for item in requested:
            if not item:
                continue
            wanted.update(self.resolve_source_ids(item))
            wanted.update(token.upper() for token in SOURCE_ID_RE.findall(item.upper()))
        unknown = sorted(wanted.difference(self._by_id))
        if unknown:
            raise ValueError(f"Неизвестные источники: {', '.join(unknown)}")
        return [source for source in self.sources if source.id in wanted]

    def validate(self) -> None:
        source_ids = [source.id for source in self.sources]
        if len(source_ids) != len(set(source_ids)):
            raise ValueError("В листе Источники есть повторяющиеся ID")
        field_ids = [field.id for field in self.fields]
        if len(field_ids) != len(set(field_ids)):
            raise ValueError("В листе Карта данных есть повторяющиеся ID")
        known = set(source_ids)
        for definition in self.fields:
            missing = set(definition.source_ids + definition.reserve_source_ids).difference(known)
            if missing:
                raise ValueError(f"Поле {definition.id} ссылается на неизвестные источники: {sorted(missing)}")

    def summary(self) -> dict[str, Any]:
        statuses = Counter(field.acquisition_status for field in self.fields)
        entities = Counter(field.entity for field in self.fields)
        return {
            "workbook": str(self.workbook_path),
            "sources": len(self.sources),
            "fields": len(self.fields),
            "pipeline_steps": len(self.pipeline),
            "risks": len(self.risks),
            "field_statuses": dict(statuses),
            "entities": dict(entities),
        }

    def to_dict(self) -> dict[str, Any]:
        return {
            "workbook": str(self.workbook_path),
            "sources": [source.to_dict() for source in self.sources],
            "fields": [definition.to_dict() for definition in self.fields],
            "pipeline": [step.to_dict() for step in self.pipeline],
            "risks": self.risks,
        }


def _sheet_key(value: str) -> str:
    return norm(value).replace(" ", "")


def _required_sheet(sheets: dict[str, Any], name: str) -> Any:
    sheet = sheets.get(_sheet_key(name))
    if sheet is None:
        raise ValueError(f"В source map отсутствует лист: {name}")
    return sheet


def resolve_source_ids(value: str | None, aliases: dict[str, str]) -> tuple[str, ...]:
    text = clean_cell(value) or ""
    if not text:
        return ()
    tokens = [token.upper() for token in SOURCE_ID_RE.findall(text.upper())]
    if tokens:
        return tuple(dict.fromkeys(tokens))
    exact = aliases.get(norm(text))
    if exact:
        return (exact,)
    matches = [source_id for alias, source_id in aliases.items() if alias and alias in norm(text)]
    return tuple(dict.fromkeys(matches))
